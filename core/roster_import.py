"""台账导入解析器 —— 从 Excel / Word / PDF 识别在管供应商清单，一键导入台账。

设计意图
--------
银行风控人员的在管供应商清单通常来自采购/合同台账（Excel 为主），
也可能以 Word/PDF 形式归档。本模块把这类文件解析成 RosterItem 列表，
供「供应商台账」页一键导入（预览后批量 upsert）。

解析约定（Excel/Word 表格，第一行为表头，列名别名容错）：
  供应商名称（必填）/ 外包分类 / 是否重要外包 / 是否非驻场 / 是否跨境
  合同开始 / 合同结束 / 最近尽调 / 最近实地检查 / 业务条线 / 集团 / 年度金额

Word/PDF 走文本启发式：识别「供应商名称：XXX」块并尽力抽取字段。
"""
import os
import re
from datetime import datetime, date

from core.roster import OUTSOURCE_CATEGORIES

# 列名 → 台账字段（别名容错）
_COL_ALIASES = {
    "name": ["供应商名称", "供应商", "名称", "公司名称", "服务商名称", "服务提供商", "企业名称"],
    "category": ["外包分类", "分类", "类别", "外包类别", "外包类型", "外包活动"],
    "important": ["重要外包", "是否重要", "是否重要外包", "重要", "重要程度"],
    "non_resident": ["非驻场", "是否非驻场", "驻场方式", "驻场"],
    "cross_border": ["跨境", "是否跨境", "跨境外包"],
    "contract_start": ["合同开始", "合同起始", "合同开始日期", "合同起始日", "合同签订"],
    "contract_end": ["合同结束", "合同到期", "合同结束日期", "合同到期日", "合同终止"],
    "last_dd": ["尽调日期", "最近尽调", "尽职调查日期", "最近尽调日期", "尽调"],
    "last_onsite": ["实地检查", "最近检查", "最近实地检查", "现场检查日期", "实地检查日期"],
    "business_lines": ["业务条线", "条线", "业务线", "涉及业务", "业务范围"],
    "group": ["集团", "关联方", "归集", "集团归集", "关联方归集"],
    "annual_amount": ["年度金额", "金额", "年金额", "年度外包金额", "外包金额", "合同金额"],
    "source": ["数据来源", "来源", "数据口径", "台账来源"],
}


def _norm_col(name: str) -> str:
    """列名 → 台账字段 key（未识别返回空串）。"""
    if not name:
        return ""
    n = str(name).strip().replace(" ", "").replace("　", "")
    for key, aliases in _COL_ALIASES.items():
        for a in aliases:
            if a.replace(" ", "") == n or n in a.replace(" ", "") or a.replace(" ", "") in n:
                return key
    return ""


def _to_bool(v) -> bool:
    if v is None:
        return False
    s = str(v).strip()
    return s in ("是", "true", "True", "1", "Y", "y", "重要", "非驻场", "跨境")


def _to_date(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if s in ("—", "-", "无", "null", "None"):
        return ""
    return re.sub(r"[./]", "-", s)


def _to_float(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(str(v).replace(",", "").replace("，", "").replace("万元", "").replace("万", "").strip())
    except (ValueError, TypeError):
        return 0.0


def _to_lines(v) -> list:
    if v in (None, ""):
        return []
    return [x.strip() for x in re.split(r"[，,、;；/|]", str(v)) if x.strip()]


def _map_category(v) -> str:
    """外包分类模糊匹配 → 标准五类（失败返回空，由前端/导入时给默认）。"""
    if v in (None, ""):
        return ""
    s = str(v).strip()
    for c in OUTSOURCE_CATEGORIES:
        core = c.replace("类", "")
        if core and (core in s or s in core):
            return c
    if "咨询" in s or "规划" in s:
        return "咨询规划类"
    if "开发" in s or "测试" in s or "研发" in s:
        return "开发测试类"
    if "运维" in s or "维护" in s or "运行" in s:
        return "运行维护类"
    if "安全" in s:
        return "安全服务类"
    if "支持" in s or "客服" in s or "作业" in s or "业务" in s:
        return "业务支持类"
    return ""


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------
def parse_roster_file(file_path: str) -> dict:
    """按扩展名分发解析，返回 {ok, rows, summary, source}。"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext in (".xlsx", ".xls"):
        rows, src = _parse_excel(file_path), "Excel"
    elif ext in (".docx", ".doc"):
        rows, src = _parse_docx(file_path), "Word"
    elif ext == ".pdf":
        rows, src = _parse_pdf(file_path), "PDF"
    else:
        raise ValueError(f"不支持的格式 {ext}，支持 Excel/Word/PDF")

    return _summarize(rows, src)


def _summarize(rows: list, src: str) -> dict:
    """统计 + 校验，返回统一结构（供预览与提交）。"""
    valid = sum(1 for r in rows if r.get("valid"))
    return {
        "ok": bool(rows),
        "source": src,
        "rows": rows,
        "summary": {
            "total": len(rows),
            "valid": valid,
            "invalid": len(rows) - valid,
        },
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _parse_excel(file_path: str) -> list:
    from openpyxl import load_workbook
    # Workbook 不支持上下文管理器协议（with 会 TypeError: ... missed __exit__ method，
    # 线上表现为「文件解析失败：__enter__」），但句柄必须显式 close()：
    # Windows 下不关闭会锁住临时文件，导致调用方的 os.unlink 失败。
    wb = load_workbook(file_path, data_only=True)
    try:
        ws = wb.active
        raw = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not raw:
        raise ValueError("Excel 为空")
    header = [str(h).strip() if h else "" for h in raw[0]]
    return _rows_from_table(header, raw[1:])


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def _parse_docx(file_path: str) -> list:
    from docx import Document
    d = Document(file_path)
    # 优先读表格（清单通常以表格呈现）
    for table in d.tables:
        data = table.rows
        if not data:
            continue
        header = [c.text.strip() for c in data[0].cells]
        # 表头含"供应商名称"类列才算清单表
        if any(_norm_col(h) == "name" for h in header):
            body = [[c.text for c in row.cells] for row in data[1:]]
            return _rows_from_table(header, body)
    # 无表格 → 段落文本启发式
    text = "\n".join(p.text for p in d.paragraphs)
    return _rows_from_text(text)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _parse_pdf(file_path: str) -> list:
    import fitz
    with fitz.open(file_path) as doc:
        text = "\n".join(page.get_text() for page in doc)
    return _rows_from_text(text)


# ---------------------------------------------------------------------------
# 表格行解析（Excel / Word 表格共用）
# ---------------------------------------------------------------------------
def _rows_from_table(header: list, data_rows: list) -> list:
    # 建立列 → 字段映射
    col_map = {}
    for i, h in enumerate(header):
        k = _norm_col(h)
        if k and k not in col_map:
            col_map[k] = i

    # 列 → 取值函数。**只处理文件里真有的列**：
    # 缺列时不能凭空填默认值，否则同名供应商导入时会把台账里已有的
    # 分类 / 合同起止 / 业务条线用 "开发测试类" / "" / [] 覆盖掉（数据丢失）。
    spec = [
        ("category", lambda v: _map_category(v) or "开发测试类"),
        ("important", _to_bool),
        ("non_resident", _to_bool),
        ("cross_border", _to_bool),
        ("contract_start", _to_date),
        ("contract_end", _to_date),
        ("last_dd", _to_date),
        ("last_onsite", _to_date),
        ("business_lines", _to_lines),
        ("group", lambda v: str(v or "").strip()),
        ("annual_amount", _to_float),
        ("source", lambda v: str(v or "").strip()),
    ]

    rows = []
    for r in data_rows:
        def cell(key):
            i = col_map.get(key)
            return r[i] if i is not None and i < len(r) else None

        name = str(cell("name") or "").strip() if cell("name") else ""
        if not name:
            continue  # 无名称的行跳过
        row = {"name": name, "valid": True}
        for key, fn in spec:
            if key in col_map:
                row[key] = fn(cell(key))
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# 文本启发式（Word 段落 / PDF）
# ---------------------------------------------------------------------------
def _rows_from_text(text: str) -> list:
    """从纯文本识别「供应商名称：XXX」块，尽力抽取字段。"""
    rows = []
    # 找到所有名称锚点
    name_pat = re.compile(r"(?:供应商名称|公司名称|企业名称|供应商|服务商)[：:\s]*([^\n，,；;。]{2,40})")
    for m in name_pat.finditer(text):
        name = m.group(1).strip()
        if not name or len(name) < 2:
            continue
        # 该名称之后 200 字作为该供应商的字段块
        seg = text[m.end(): m.end() + 200]
        def grab(*keys):
            for k in keys:
                mm = re.search(k + r"[：:]\s*([^\n，,；;。]{1,40})", seg)
                if mm:
                    return mm.group(1).strip()
            return ""
        category = _map_category(grab("外包分类", "分类"))
        amt = _to_float(grab("年度金额", "金额"))
        rows.append({
            "name": name,
            "category": category or "开发测试类",
            "important": _to_bool(grab("重要外包", "是否重要")),
            "non_resident": _to_bool(grab("驻场方式", "是否非驻场")),
            "cross_border": _to_bool(grab("是否跨境")),
            "contract_start": grab("合同开始"),
            "contract_end": grab("合同结束"),
            "last_dd": grab("尽调日期", "最近尽调"),
            "last_onsite": "",
            "business_lines": _to_lines(grab("业务条线", "条线")),
            "group": grab("集团", "关联方"),
            "annual_amount": amt,
            "valid": True,
        })
    return rows
